import json
import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from pathlib import Path
from typing import Any

ROADMAP_FILE = Path("RoadMap_Shipment.json")
GREENMILE_FILE = Path("GreenMile_Estados.json")
OUTPUT_FILE = Path("PowerApp_Data.json")
OUTPUT_XLSX = Path("PowerApp_Data.xlsx")


def load_json_array(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo requerido: {path}")

    with path.open("r", encoding="utf-8") as file:
        data = json.load(file)

    if not isinstance(data, list):
        raise ValueError(f"{path} no contiene una matriz JSON.")

    return data


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def optional_number(value: Any) -> Any:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def main() -> None:
    roadmap_records = load_json_array(ROADMAP_FILE)
    greenmile_records = load_json_array(GREENMILE_FILE)

    greenmile_by_shipment: dict[str, dict[str, Any]] = {}

    for position, record in enumerate(greenmile_records, start=1):
        if not isinstance(record, dict):
            raise ValueError(
                f"Registro {position} de {GREENMILE_FILE} no es un objeto."
            )

        shipment = clean_text(record.get("ShipmentCustom"))
        if shipment == "":
            continue

        greenmile_by_shipment[shipment] = record

    output_records: list[dict[str, Any]] = []
    seen_shipments: set[str] = set()

    for position, roadmap in enumerate(roadmap_records, start=1):
        if not isinstance(roadmap, dict):
            raise ValueError(
                f"Registro {position} de {ROADMAP_FILE} no es un objeto."
            )

        shipment = clean_text(roadmap.get("ShipmentCustom"))
        if shipment == "":
            continue

        if shipment in seen_shipments:
            raise ValueError(
                f"ShipmentCustom duplicado en RoadMap_Shipment.json: {shipment}"
            )
        seen_shipments.add(shipment)

        greenmile = greenmile_by_shipment.get(shipment, {})

        output_record = {
            "DeliveryDate": clean_text(roadmap.get("DeliveryDate")),
            "ShipmentCustom": shipment,
            "VehicleKey": clean_text(roadmap.get("VehicleKey")),
            "Location": clean_text(roadmap.get("Location")),
            "CustomerAccount": clean_text(roadmap.get("CustomerAccount")),
            "CustomerName": clean_text(roadmap.get("CustomerName")),
            "DriverName": clean_text(roadmap.get("DriverName")),
            "DriverBadge": clean_text(roadmap.get("DriverBadge")),
            "RutaVenta": clean_text(roadmap.get("RutaVenta")),
            "Canal": clean_text(roadmap.get("Canal")),
            "KgPlanificados": optional_number(
                roadmap.get("KgPlanificados")
            ),
            "EstadoMigracion": clean_text(
                greenmile.get("EstadoMigracion") or "NO MIGRADO"
            ),
            "EstadoConexion": clean_text(
                greenmile.get("EstadoConexion") or "NO APLICA"
            ),
            "EstadoEntrega": clean_text(
                greenmile.get("EstadoEntrega") or "SIN INFORMACION"
            ),
            "Latitude": optional_number(
                roadmap.get("Latitude")
            ),
            "Longitude": optional_number(
                roadmap.get("Longitude")
            ),
            "LatitudActual": optional_number(
                greenmile.get("LatitudActual")
            ),
            "LongitudActual": optional_number(
                greenmile.get("LongitudActual")
            ),
            "FuenteCoordenada": clean_text(
                greenmile.get("FuenteCoordenada")
            ),
            "MotivoNoEntrega": clean_text(
                greenmile.get("MotivoNoEntrega")
            ),
            "MotivoCancelacion": clean_text(
                greenmile.get("MotivoCancelacion")
            ),
        }
        output_records.append(output_record)

    if len(output_records) == 0:
        raise ValueError("No se generaron registros para Power Apps.")

    with OUTPUT_FILE.open("w", encoding="utf-8", newline="") as file:
        json.dump(output_records, file, ensure_ascii=False, indent=2)
        file.write("\n")

        # Generar Excel para Power Apps
    dataframe = pd.DataFrame(output_records)

    with pd.ExcelWriter(
        OUTPUT_XLSX,
        engine="openpyxl"
    ) as writer:
        dataframe.to_excel(
            writer,
            index=False,
            sheet_name="Data"
        )

    workbook = load_workbook(OUTPUT_XLSX)
    worksheet = workbook["Data"]

    last_row = worksheet.max_row
    last_col = worksheet.max_column

    table = Table(
        displayName="PowerAppData",
        ref=f"A1:{chr(64 + last_col)}{last_row}"
    )

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True
    )

    table.tableStyleInfo = style
    worksheet.add_table(table)

    workbook.save(OUTPUT_XLSX)
    
    migrated = sum(
        record["EstadoMigracion"] == "MIGRADO"
        for record in output_records
    )
    connected = sum(
        record["EstadoConexion"] == "CONECTADO"
        for record in output_records
    )
    delivered = sum(
        record["EstadoEntrega"] == "ENTREGADO"
        for record in output_records
    )
    with_coordinates = sum(
        record["LatitudActual"] is not None
        and record["LongitudActual"] is not None
        for record in output_records
    )

    print(f"Registros Power Apps: {len(output_records)}")
    print(f"Shipments migrados: {migrated}")
    print(f"Shipments conectados: {connected}")
    print(f"Shipments entregados: {delivered}")
    print(f"Shipments con coordenadas operativas: {with_coordinates}")
    print(f"Archivo generado: {OUTPUT_FILE}")
    print(f"Archivo generado: {OUTPUT_XLSX}")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise
